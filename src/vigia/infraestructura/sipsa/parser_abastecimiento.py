"""
Parser de microdatos de abastecimiento SIPSA.

Convierte las hojas de microdatos publicadas por DANE en registros canónicos
de abastecimiento.

El parser acepta tanto contenido XLSX en memoria como una ruta local. Para
archivos grandes debe utilizarse una ruta, permitiendo que openpyxl opere en
modo de solo lectura sin cargar previamente el archivo completo en RAM.
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from vigia.dominio.sipsa import RegistroAbastecimiento


class ErrorParserAbastecimientoSipsa(RuntimeError):
    """Error producido al interpretar microdatos SIPSA de abastecimiento."""


HOJAS_MICRODATOS = ("2.1", "2.2", "2.3")

ENCABEZADOS_ESPERADOS = (
    "Ciudad, Mercado Mayorista",
    "Fecha",
    "Divipola Depto Proc.",
    "Divipola Municipio / ISO 3166-1 País Proc.",
    "Departamento Proc.",
    "Municipio de Colombia / País Proc.",
    "Grupo",
    "Código CPC",
    "Alimento",
    "Cant Kg",
)

PREFIJOS_PIE_DOCUMENTAL = (
    "Fuente:",
    "Fecha de actualización:",
)

FIRMA_XLSX = b"PK\x03\x04"


def _texto_obligatorio(valor: object, campo: str) -> str:
    """Convierte y valida un campo textual obligatorio."""
    if valor is None:
        raise ErrorParserAbastecimientoSipsa(f"{campo} no puede estar vacío.")

    texto = str(valor).strip()

    if not texto:
        raise ErrorParserAbastecimientoSipsa(f"{campo} no puede estar vacío.")

    return texto


def _normalizar_codigo(valor: object, campo: str) -> str:
    """
    Normaliza códigos importados desde Excel.

    SIPSA publica algunos códigos como cadenas precedidas por apóstrofo.
    Se conservan como texto para evitar pérdida de ceros iniciales.
    """
    texto = _texto_obligatorio(valor, campo)

    if texto.startswith("'"):
        texto = texto[1:]

    if not texto:
        raise ErrorParserAbastecimientoSipsa(f"{campo} quedó vacío después de normalizarse.")

    return texto


def _convertir_fecha(valor: object) -> date:
    """Convierte la representación temporal SIPSA a fecha."""
    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    if isinstance(valor, str):
        texto = valor.strip()

        for formato in (
            "%Y-%m-%d",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                return datetime.strptime(
                    texto,
                    formato,
                ).date()
            except ValueError:
                continue

    raise ErrorParserAbastecimientoSipsa(f"Fecha SIPSA inválida: {valor!r}")


def _convertir_cantidad(valor: object) -> Decimal:
    """Convierte la cantidad de abastecimiento a Decimal."""
    if valor is None:
        raise ErrorParserAbastecimientoSipsa("Cant Kg no puede estar vacía.")

    try:
        cantidad = Decimal(str(valor))
    except (InvalidOperation, ValueError) as exc:
        raise ErrorParserAbastecimientoSipsa(f"Cantidad SIPSA inválida: {valor!r}") from exc

    if cantidad < 0:
        raise ErrorParserAbastecimientoSipsa(f"Cantidad SIPSA negativa: {cantidad}")

    return cantidad


def _validar_encabezados(
    encabezados: tuple[object, ...],
    *,
    hoja: str,
) -> None:
    """Comprueba que el esquema DANE siga siendo compatible."""
    reales = tuple(
        str(valor).strip() if valor is not None else ""
        for valor in encabezados[: len(ENCABEZADOS_ESPERADOS)]
    )

    if reales != ENCABEZADOS_ESPERADOS:
        raise ErrorParserAbastecimientoSipsa(
            f"Esquema inesperado en hoja {hoja!r}. "
            f"Esperado: {ENCABEZADOS_ESPERADOS!r}. "
            f"Recibido: {reales!r}."
        )


def _es_fila_vacia(
    valores: tuple[object, ...],
) -> bool:
    """Indica si una fila no contiene valores de microdatos."""
    return all(valor is None for valor in valores)


def _es_pie_documental(
    valores: tuple[object, ...],
) -> bool:
    """Detecta líneas editoriales agregadas por DANE al final de una hoja."""
    primer_valor = valores[0]

    if not isinstance(primer_valor, str):
        return False

    texto = primer_valor.strip()

    return any(texto.startswith(prefijo) for prefijo in PREFIJOS_PIE_DOCUMENTAL)


def _validar_archivo_local(ruta: Path) -> None:
    """Valida existencia, tipo y firma básica de un XLSX local."""
    if not ruta.exists():
        raise ErrorParserAbastecimientoSipsa(f"El archivo SIPSA no existe: {ruta}")

    if not ruta.is_file():
        raise ErrorParserAbastecimientoSipsa(f"La ruta SIPSA no corresponde a un archivo: {ruta}")

    try:
        with ruta.open("rb") as archivo:
            firma = archivo.read(4)
    except OSError as exc:
        raise ErrorParserAbastecimientoSipsa(
            f"No fue posible leer el archivo SIPSA: {ruta}"
        ) from exc

    if firma != FIRMA_XLSX:
        raise ErrorParserAbastecimientoSipsa("El contenido recibido no parece un XLSX válido.")


def _preparar_fuente(
    contenido: bytes | Path,
) -> BinaryIO | Path:
    """
    Prepara la fuente para openpyxl.

    Los bytes se mantienen por compatibilidad con pruebas y archivos pequeños.
    Las rutas se entregan directamente a openpyxl para evitar duplicar archivos
    grandes completos en memoria.
    """
    if isinstance(contenido, Path):
        _validar_archivo_local(contenido)
        return contenido

    if not contenido.startswith(FIRMA_XLSX):
        raise ErrorParserAbastecimientoSipsa("El contenido recibido no parece un XLSX válido.")

    return BytesIO(contenido)


def parsear_abastecimiento_sipsa(
    contenido: bytes | Path,
    *,
    archivo_fuente: str,
) -> Iterator[RegistroAbastecimiento]:
    """
    Genera registros canónicos desde microdatos XLSX de abastecimiento SIPSA.

    Para archivos grandes se recomienda suministrar ``Path``. Las hojas se
    recorren secuencialmente con ``read_only=True`` y cada registro se entrega
    mediante un iterador, manteniendo el consumo de memoria acotado.
    """
    fuente = _preparar_fuente(contenido)

    try:
        libro = load_workbook(
            fuente,
            read_only=True,
            data_only=True,
        )
    except (OSError, ValueError) as exc:
        raise ErrorParserAbastecimientoSipsa(
            "No fue posible abrir el archivo XLSX de abastecimiento."
        ) from exc

    try:
        hojas_disponibles = [nombre for nombre in HOJAS_MICRODATOS if nombre in libro.sheetnames]

        if not hojas_disponibles:
            raise ErrorParserAbastecimientoSipsa("No se encontraron hojas de microdatos SIPSA.")

        for nombre_hoja in hojas_disponibles:
            hoja = libro[nombre_hoja]

            iterador = hoja.iter_rows(
                min_row=9,
                values_only=True,
            )

            try:
                encabezados = next(iterador)
            except StopIteration as exc:
                raise ErrorParserAbastecimientoSipsa(
                    f"La hoja {nombre_hoja!r} está vacía."
                ) from exc

            _validar_encabezados(
                encabezados,
                hoja=nombre_hoja,
            )

            for numero_fila, fila in enumerate(
                iterador,
                start=10,
            ):
                valores = fila[:10]

                if _es_fila_vacia(valores):
                    continue

                if _es_pie_documental(valores):
                    break

                try:
                    yield RegistroAbastecimiento(
                        fecha=_convertir_fecha(valores[1]),
                        ciudad_mercado=_texto_obligatorio(
                            valores[0],
                            "Ciudad, Mercado Mayorista",
                        ),
                        codigo_departamento_origen=_normalizar_codigo(
                            valores[2],
                            "Divipola Depto Proc.",
                        ),
                        codigo_municipio_pais_origen=_normalizar_codigo(
                            valores[3],
                            "Divipola Municipio / ISO 3166-1 País Proc.",
                        ),
                        departamento_origen=_texto_obligatorio(
                            valores[4],
                            "Departamento Proc.",
                        ),
                        municipio_pais_origen=_texto_obligatorio(
                            valores[5],
                            "Municipio de Colombia / País Proc.",
                        ),
                        grupo=_texto_obligatorio(
                            valores[6],
                            "Grupo",
                        ),
                        codigo_cpc=_normalizar_codigo(
                            valores[7],
                            "Código CPC",
                        ),
                        producto=_texto_obligatorio(
                            valores[8],
                            "Alimento",
                        ),
                        cantidad_kg=_convertir_cantidad(valores[9]),
                        fuente="SIPSA-DANE",
                        archivo_fuente=archivo_fuente,
                        hoja_fuente=nombre_hoja,
                    )

                except (
                    ErrorParserAbastecimientoSipsa,
                    ValueError,
                ) as exc:
                    raise ErrorParserAbastecimientoSipsa(
                        f"Registro inválido en hoja {nombre_hoja!r}, fila {numero_fila}: {exc}"
                    ) from exc

    finally:
        libro.close()
