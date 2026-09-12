"""
Parser de precios mayoristas SIPSA.

Convierte la estructura matricial del boletín diario de precios del DANE
en registros canónicos de dominio.

El parser acepta contenido XLSX en memoria o una ruta local y recorre la
hoja secuencialmente para evitar accesos aleatorios costosos en modo
read_only de openpyxl.
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from vigia.dominio.sipsa import RegistroPrecioMayorista


class ErrorParserPreciosSipsa(RuntimeError):
    """Error producido al interpretar un archivo SIPSA de precios."""


FIRMA_XLSX = b"PK\x03\x04"


def _convertir_decimal(valor: object) -> Decimal | None:
    """Convierte valores numéricos SIPSA a Decimal."""
    if valor is None:
        return None

    if isinstance(valor, str):
        texto = valor.strip()

        if not texto or texto.lower() == "n.d.":
            return None

        texto = texto.replace(",", ".")

        try:
            return Decimal(texto)
        except InvalidOperation as exc:
            raise ErrorParserPreciosSipsa(f"Valor numérico inválido en SIPSA: {valor!r}") from exc

    try:
        return Decimal(str(valor))
    except (InvalidOperation, ValueError) as exc:
        raise ErrorParserPreciosSipsa(f"Valor numérico inválido en SIPSA: {valor!r}") from exc


def _extraer_fecha(valor: object) -> datetime:
    """Convierte el encabezado de fecha del boletín SIPSA a datetime."""
    if not isinstance(valor, str):
        raise ErrorParserPreciosSipsa(f"Fecha SIPSA inesperada: {valor!r}")

    texto = valor.strip().lower()

    meses = {
        "enero": "January",
        "febrero": "February",
        "marzo": "March",
        "abril": "April",
        "mayo": "May",
        "junio": "June",
        "julio": "July",
        "agosto": "August",
        "septiembre": "September",
        "octubre": "October",
        "noviembre": "November",
        "diciembre": "December",
    }

    dias = {
        "lunes": "Monday",
        "martes": "Tuesday",
        "miércoles": "Wednesday",
        "miercoles": "Wednesday",
        "jueves": "Thursday",
        "viernes": "Friday",
        "sábado": "Saturday",
        "sabado": "Saturday",
        "domingo": "Sunday",
    }

    for espanol, ingles in dias.items():
        texto = texto.replace(espanol, ingles)

    for espanol, ingles in meses.items():
        texto = texto.replace(espanol, ingles)

    try:
        return datetime.strptime(
            texto,
            "%A %d de %B de %Y",
        )
    except ValueError as exc:
        raise ErrorParserPreciosSipsa(
            f"No fue posible interpretar la fecha SIPSA: {valor!r}"
        ) from exc


def _separar_ciudad_mercado(
    valor: object,
) -> tuple[str, str]:
    """Interpreta el encabezado territorial publicado por SIPSA."""
    if not isinstance(valor, str):
        raise ErrorParserPreciosSipsa(f"Mercado SIPSA inválido: {valor!r}")

    texto = valor.strip()

    if not texto:
        raise ErrorParserPreciosSipsa("El encabezado de mercado SIPSA está vacío.")

    if "," not in texto:
        return texto, texto

    ciudad, mercado = (parte.strip() for parte in texto.split(",", maxsplit=1))

    if not ciudad or not mercado:
        raise ErrorParserPreciosSipsa(f"Encabezado de mercado SIPSA inválido: {valor!r}")

    return ciudad, mercado


def _validar_archivo_local(ruta: Path) -> None:
    """Valida existencia y firma básica del XLSX."""
    if not ruta.exists():
        raise ErrorParserPreciosSipsa(f"El archivo SIPSA no existe: {ruta}")

    if not ruta.is_file():
        raise ErrorParserPreciosSipsa(f"La ruta SIPSA no corresponde a un archivo: {ruta}")

    try:
        with ruta.open("rb") as archivo:
            firma = archivo.read(4)
    except OSError as exc:
        raise ErrorParserPreciosSipsa(f"No fue posible leer el archivo SIPSA: {ruta}") from exc

    if firma != FIRMA_XLSX:
        raise ErrorParserPreciosSipsa("El contenido recibido no parece un XLSX válido.")


def _preparar_fuente(
    contenido: bytes | Path,
) -> BinaryIO | Path:
    """Prepara una fuente XLSX para openpyxl."""
    if isinstance(contenido, Path):
        _validar_archivo_local(contenido)
        return contenido

    if not contenido.startswith(FIRMA_XLSX):
        raise ErrorParserPreciosSipsa("El contenido recibido no parece un XLSX válido.")

    return BytesIO(contenido)


def parsear_precios_sipsa(
    contenido: bytes | Path,
    *,
    archivo_fuente: str,
) -> Iterator[RegistroPrecioMayorista]:
    """
    Genera registros canónicos desde un boletín diario SIPSA.

    La hoja se recorre una única vez de manera secuencial.
    """
    fuente = _preparar_fuente(contenido)

    try:
        libro = load_workbook(
            fuente,
            read_only=True,
            data_only=True,
        )
    except (OSError, ValueError) as exc:
        raise ErrorParserPreciosSipsa("No fue posible abrir el archivo XLSX de precios.") from exc

    try:
        if "Boletín diario" not in libro.sheetnames:
            raise ErrorParserPreciosSipsa("No existe la hoja esperada 'Boletín diario'.")

        hoja = libro["Boletín diario"]

        fecha = None
        mercados: dict[int, tuple[str, str]] = {}
        categoria_actual: str | None = None

        for numero_fila, fila in enumerate(
            hoja.iter_rows(values_only=True),
            start=1,
        ):
            if numero_fila == 2:
                fecha = _extraer_fecha(fila[0]).date()
                continue

            if numero_fila == 3:
                for indice in range(1, len(fila), 2):
                    valor_mercado = fila[indice]

                    if valor_mercado is None:
                        continue

                    mercados[indice] = _separar_ciudad_mercado(valor_mercado)

                if not mercados:
                    raise ErrorParserPreciosSipsa("No se encontraron mercados en el boletín SIPSA.")

                continue

            if numero_fila < 5:
                continue

            producto = fila[0]

            if producto is None:
                continue

            producto_texto = str(producto).strip()

            if not producto_texto:
                continue

            valores_mercado = [fila[indice] if indice < len(fila) else None for indice in mercados]

            if all(valor is None for valor in valores_mercado):
                categoria_actual = producto_texto
                continue

            if fecha is None:
                raise ErrorParserPreciosSipsa("No se encontró la fecha del boletín SIPSA.")

            if categoria_actual is None:
                raise ErrorParserPreciosSipsa(
                    f"Producto sin categoría en fila {numero_fila}: {producto_texto}"
                )

            for indice, (ciudad, mercado) in mercados.items():
                precio = _convertir_decimal(fila[indice] if indice < len(fila) else None)

                indice_variacion = indice + 1

                variacion = _convertir_decimal(
                    fila[indice_variacion] if indice_variacion < len(fila) else None
                )

                yield RegistroPrecioMayorista(
                    fecha=fecha,
                    producto=producto_texto,
                    categoria=categoria_actual,
                    ciudad=ciudad,
                    mercado=mercado,
                    precio_kg=precio,
                    variacion=variacion,
                    fuente="SIPSA-DANE",
                    archivo_fuente=archivo_fuente,
                )

    finally:
        libro.close()
