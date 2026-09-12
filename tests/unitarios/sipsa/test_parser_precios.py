"""
Pruebas unitarias del parser de precios mayoristas SIPSA.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook

from vigia.infraestructura.sipsa import (
    ErrorParserPreciosSipsa,
    parsear_precios_sipsa,
)


def _crear_xlsx_prueba() -> bytes:
    """Construye un boletín SIPSA mínimo con estructuras observadas realmente."""
    libro = Workbook()
    hoja = libro.active

    assert hoja is not None

    hoja.title = "Boletín diario"

    hoja.cell(
        row=2,
        column=1,
        value="Martes 1 de septiembre de 2026",
    )

    hoja.cell(row=3, column=1, value="Precio $/Kg")
    hoja.cell(row=3, column=2, value="Medellín, CMA")
    hoja.cell(row=3, column=4, value="Bogotá, Corabastos")
    hoja.cell(row=3, column=6, value="Santa Marta")

    hoja.cell(row=4, column=2, value="Precio")
    hoja.cell(row=4, column=3, value="Var %")
    hoja.cell(row=4, column=4, value="Precio")
    hoja.cell(row=4, column=5, value="Var %")
    hoja.cell(row=4, column=6, value="Precio")
    hoja.cell(row=4, column=7, value="Var %")

    hoja.cell(
        row=5,
        column=1,
        value="Verduras y hortalizas",
    )

    hoja.cell(row=6, column=1, value="Ahuyama")
    hoja.cell(row=6, column=2, value=1700)
    hoja.cell(row=6, column=3, value=0.10)
    hoja.cell(row=6, column=4, value=2350)
    hoja.cell(row=6, column=5, value=-0.01)
    hoja.cell(row=6, column=6, value=2100)
    hoja.cell(row=6, column=7, value=0.04)

    hoja.cell(row=7, column=1, value="Habichuela")
    hoja.cell(row=7, column=2, value="n.d.")
    hoja.cell(row=7, column=3, value="n.d.")
    hoja.cell(row=7, column=4, value=5708)
    hoja.cell(row=7, column=5, value=-0.13)
    hoja.cell(row=7, column=6, value="n.d.")
    hoja.cell(row=7, column=7, value="n.d.")

    buffer = BytesIO()
    libro.save(buffer)
    libro.close()

    return buffer.getvalue()


def test_parsea_boletin_diario() -> None:
    """Convierte correctamente la matriz SIPSA en registros largos."""
    registros = list(
        parsear_precios_sipsa(
            _crear_xlsx_prueba(),
            archivo_fuente="prueba.xlsx",
        )
    )

    assert len(registros) == 6

    primero = registros[0]

    assert primero.producto == "Ahuyama"
    assert primero.categoria == "Verduras y hortalizas"
    assert primero.ciudad == "Medellín"
    assert primero.mercado == "CMA"
    assert str(primero.precio_kg) == "1700"
    assert str(primero.variacion) == "0.1"
    assert primero.fecha.isoformat() == "2026-09-01"


def test_preserva_valores_no_disponibles() -> None:
    """Convierte n.d. en ausencia explícita y nunca en cero."""
    registros = list(
        parsear_precios_sipsa(
            _crear_xlsx_prueba(),
            archivo_fuente="prueba.xlsx",
        )
    )

    registro = next(
        item for item in registros if item.producto == "Habichuela" and item.ciudad == "Medellín"
    )

    assert registro.precio_kg is None
    assert registro.variacion is None


def test_admite_mercado_sin_separador() -> None:
    """Conserva denominaciones SIPSA que no incluyen ciudad y mercado separados."""
    registros = list(
        parsear_precios_sipsa(
            _crear_xlsx_prueba(),
            archivo_fuente="prueba.xlsx",
        )
    )

    registro = next(
        item for item in registros if item.producto == "Ahuyama" and item.ciudad == "Santa Marta"
    )

    assert registro.ciudad == "Santa Marta"
    assert registro.mercado == "Santa Marta"
    assert str(registro.precio_kg) == "2100"


def test_rechaza_contenido_no_xlsx() -> None:
    """Rechaza entradas que no representan un libro XLSX."""
    with pytest.raises(
        ErrorParserPreciosSipsa,
        match="no parece un XLSX válido",
    ):
        list(
            parsear_precios_sipsa(
                b"contenido-invalido",
                archivo_fuente="invalido.xlsx",
            )
        )


def test_rechaza_libro_sin_hoja_esperada() -> None:
    """Falla explícitamente ante cambios estructurales del proveedor."""
    libro = Workbook()
    hoja = libro.active

    assert hoja is not None

    hoja.title = "Otra hoja"

    buffer = BytesIO()
    libro.save(buffer)
    libro.close()

    with pytest.raises(
        ErrorParserPreciosSipsa,
        match="Boletín diario",
    ):
        list(
            parsear_precios_sipsa(
                buffer.getvalue(),
                archivo_fuente="invalido.xlsx",
            )
        )
