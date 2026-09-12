"""
Pruebas unitarias del parser de abastecimiento SIPSA.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from vigia.infraestructura.sipsa import (
    ErrorParserAbastecimientoSipsa,
    parsear_abastecimiento_sipsa,
)

ENCABEZADOS = (
    "Ciudad, Mercado Mayorista",
    "Fecha",
    "Divipola Depto Proc.",
    "Divipola Municipio / ISO 3166-1 País Proc.",
    "Departamento Proc.",
    "Municipio de Colombia / País Proc.",
    "Grupo",
    "Código CPC",
    "Alimento",
    "Cant Kg",
)


def _agregar_encabezados(hoja: Worksheet) -> None:
    """Agrega los encabezados oficiales de microdatos a una hoja."""
    for columna, encabezado in enumerate(
        ENCABEZADOS,
        start=1,
    ):
        hoja.cell(
            row=9,
            column=columna,
            value=encabezado,
        )


def _agregar_registro_valido(
    hoja: Worksheet,
    *,
    fila: int = 10,
    cantidad_kg: int = 8000,
) -> None:
    """Agrega un registro SIPSA válido a una hoja de prueba."""
    hoja.cell(row=fila, column=1, value="Armenia, Mercar")
    hoja.cell(row=fila, column=2, value="2026-01-02 00:00:00")
    hoja.cell(row=fila, column=3, value="'52")
    hoja.cell(row=fila, column=4, value="'52838")
    hoja.cell(row=fila, column=5, value="NARIÑO")
    hoja.cell(row=fila, column=6, value="TÚQUERRES")
    hoja.cell(
        row=fila,
        column=7,
        value="TUBERCULOS, RAICES Y PLATANOS",
    )
    hoja.cell(row=fila, column=8, value="'0151001")
    hoja.cell(row=fila, column=9, value="Papa capira")
    hoja.cell(row=fila, column=10, value=cantidad_kg)


def _crear_xlsx_prueba() -> bytes:
    """Construye un microdato SIPSA mínimo con estructura real."""
    libro = Workbook()

    hoja_inicial = libro.active
    assert hoja_inicial is not None
    hoja_inicial.title = "Índice"

    hoja = libro.create_sheet("2.1")

    _agregar_encabezados(hoja)
    _agregar_registro_valido(hoja)

    hoja.cell(
        row=12,
        column=1,
        value="Fuente: DANE-SIPSA.",
    )
    hoja.cell(
        row=13,
        column=1,
        value="Fecha de actualización: 15 de mayo de 2026",
    )

    buffer = BytesIO()
    libro.save(buffer)
    libro.close()

    return buffer.getvalue()


def test_parsea_abastecimiento() -> None:
    """Convierte una fila SIPSA en un registro canónico."""
    registros = list(
        parsear_abastecimiento_sipsa(
            _crear_xlsx_prueba(),
            archivo_fuente="abastecimiento.xlsx",
        )
    )

    assert len(registros) == 1

    registro = registros[0]

    assert registro.fecha.isoformat() == "2026-01-02"
    assert registro.ciudad_mercado == "Armenia, Mercar"
    assert registro.codigo_departamento_origen == "52"
    assert registro.codigo_municipio_pais_origen == "52838"
    assert registro.codigo_cpc == "0151001"
    assert registro.producto == "Papa capira"
    assert str(registro.cantidad_kg) == "8000"
    assert registro.hoja_fuente == "2.1"


def test_ignora_pie_documental() -> None:
    """No interpreta como microdatos el pie editorial publicado por DANE."""
    registros = list(
        parsear_abastecimiento_sipsa(
            _crear_xlsx_prueba(),
            archivo_fuente="abastecimiento.xlsx",
        )
    )

    assert len(registros) == 1
    assert registros[0].producto == "Papa capira"


def test_rechaza_esquema_modificado() -> None:
    """Detecta cambios incompatibles en la estructura SIPSA."""
    libro = Workbook()

    hoja = libro.active
    assert hoja is not None

    hoja.title = "2.1"
    hoja.cell(
        row=9,
        column=1,
        value="Columna inesperada",
    )

    buffer = BytesIO()
    libro.save(buffer)
    libro.close()

    with pytest.raises(
        ErrorParserAbastecimientoSipsa,
        match="Esquema inesperado",
    ):
        list(
            parsear_abastecimiento_sipsa(
                buffer.getvalue(),
                archivo_fuente="invalido.xlsx",
            )
        )


def test_rechaza_cantidad_negativa() -> None:
    """Rechaza explícitamente una cantidad de abastecimiento negativa."""
    libro = Workbook()

    hoja = libro.active
    assert hoja is not None
    hoja.title = "2.1"

    _agregar_encabezados(hoja)
    _agregar_registro_valido(
        hoja,
        cantidad_kg=-100,
    )

    buffer = BytesIO()
    libro.save(buffer)
    libro.close()

    with pytest.raises(
        ErrorParserAbastecimientoSipsa,
        match="Cantidad SIPSA negativa",
    ):
        list(
            parsear_abastecimiento_sipsa(
                buffer.getvalue(),
                archivo_fuente="abastecimiento.xlsx",
            )
        )


def test_rechaza_contenido_no_xlsx() -> None:
    """Rechaza contenido que no sea un XLSX."""
    with pytest.raises(
        ErrorParserAbastecimientoSipsa,
        match="no parece un XLSX válido",
    ):
        list(
            parsear_abastecimiento_sipsa(
                b"no-es-xlsx",
                archivo_fuente="archivo.xlsx",
            )
        )
