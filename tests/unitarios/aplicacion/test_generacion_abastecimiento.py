"""
Pruebas del caso de uso de generación diaria SIPSA_A.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from vigia.aplicacion.series_abastecimiento import (
    ErrorGeneracionAbastecimiento,
    generar_serie_abastecimiento_sipsa,
)

_ENCABEZADOS = (
    "Ciudad, Mercado Mayorista",
    "Fecha",
    "Divipola Depto Proc.",
    "Divipola Municipio / ISO 3166-1 País Proc.",
    "Departamento Proc.",
    "Municipio de Colombia / País Proc.",
    "Grupo",
    "Código CPC",
    "Alimento",
    "Cant Kg",
)


def _crear_xlsx(
    ruta: Path,
) -> None:
    """Crea una fuente SIPSA_A mínima compatible con el parser real."""
    libro = Workbook()

    hoja = libro.active

    if hoja is None:
        raise RuntimeError("openpyxl no creó la hoja inicial esperada.")

    hoja.title = "2.1"

    for fila in range(
        1,
        9,
    ):
        hoja.cell(
            row=fila,
            column=1,
            value=f"Metadato {fila}",
        )

    for columna, encabezado in enumerate(
        _ENCABEZADOS,
        start=1,
    ):
        hoja.cell(
            row=9,
            column=columna,
            value=encabezado,
        )

    hoja.append(
        (
            "Manizales, Centro Galerías",
            "2026-01-02",
            "'17",
            "'17001",
            "Caldas",
            "Manizales",
            "Tubérculos, raíces y plátanos",
            "'01234",
            "Plátano hartón verde",
            100,
        )
    )

    hoja.append(
        (
            "Manizales, Centro Galerías",
            "2026-01-02",
            "'17",
            "'17001",
            "Caldas",
            "Manizales",
            "Tubérculos, raíces y plátanos",
            "'01234",
            "Plátano hartón verde",
            150,
        )
    )

    hoja.append(
        (
            "Manizales, Centro Galerías",
            "2026-01-03",
            "'17",
            "'17001",
            "Caldas",
            "Manizales",
            "Tubérculos, raíces y plátanos",
            "'01234",
            "Plátano hartón verde",
            200,
        )
    )

    ruta.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    libro.save(ruta)

    libro.close()


def test_genera_serie_abastecimiento_completa(
    tmp_path: Path,
) -> None:
    fuente = tmp_path / "abastecimiento.xlsx"

    temporal = tmp_path / "agregacion.sqlite"

    salida = tmp_path / "salida"

    _crear_xlsx(fuente)

    resultado = generar_serie_abastecimiento_sipsa(
        ruta_fuente=fuente,
        directorio_salida=salida,
        ruta_bd_temporal=temporal,
    )

    assert resultado.archivo_fuente == "abastecimiento.xlsx"

    persistencia = resultado.persistencia

    assert persistencia.observaciones == 2
    assert persistencia.registros_fuente == 3
    assert str(persistencia.cantidad_total_kg) == "450"

    assert persistencia.fecha_minima == "2026-01-02"
    assert persistencia.fecha_maxima == "2026-01-03"

    assert persistencia.productos == 1
    assert persistencia.mercados == 1
    assert persistencia.cpc == 1
    assert persistencia.origenes == 1

    assert persistencia.ruta_datos.exists()
    assert persistencia.ruta_manifiesto.exists()

    assert not temporal.exists()


def test_resultado_es_reproducible(
    tmp_path: Path,
) -> None:
    fuente = tmp_path / "abastecimiento.xlsx"

    _crear_xlsx(fuente)

    resultado_a = generar_serie_abastecimiento_sipsa(
        ruta_fuente=fuente,
        directorio_salida=tmp_path / "a",
        ruta_bd_temporal=tmp_path / "a.sqlite",
    )

    resultado_b = generar_serie_abastecimiento_sipsa(
        ruta_fuente=fuente,
        directorio_salida=tmp_path / "b",
        ruta_bd_temporal=tmp_path / "b.sqlite",
    )

    assert resultado_a.persistencia.sha256_datos == resultado_b.persistencia.sha256_datos

    assert (
        resultado_a.persistencia.ruta_datos.read_bytes()
        == resultado_b.persistencia.ruta_datos.read_bytes()
    )

    assert (
        resultado_a.persistencia.ruta_manifiesto.read_bytes()
        == resultado_b.persistencia.ruta_manifiesto.read_bytes()
    )


def test_rechaza_fuente_inexistente(
    tmp_path: Path,
) -> None:
    with pytest.raises(
        ErrorGeneracionAbastecimiento,
        match="No existe",
    ):
        generar_serie_abastecimiento_sipsa(
            ruta_fuente=tmp_path / "inexistente.xlsx",
            directorio_salida=tmp_path / "salida",
            ruta_bd_temporal=tmp_path / "temporal.sqlite",
        )


def test_rechaza_fuente_que_no_es_archivo(
    tmp_path: Path,
) -> None:
    fuente = tmp_path / "directorio"

    fuente.mkdir()

    with pytest.raises(
        ErrorGeneracionAbastecimiento,
        match="no es un archivo",
    ):
        generar_serie_abastecimiento_sipsa(
            ruta_fuente=fuente,
            directorio_salida=tmp_path / "salida",
            ruta_bd_temporal=tmp_path / "temporal.sqlite",
        )


def test_rechaza_temporal_sobre_fuente(
    tmp_path: Path,
) -> None:
    fuente = tmp_path / "abastecimiento.xlsx"

    _crear_xlsx(fuente)

    with pytest.raises(
        ErrorGeneracionAbastecimiento,
        match="misma ruta",
    ):
        generar_serie_abastecimiento_sipsa(
            ruta_fuente=fuente,
            directorio_salida=tmp_path / "salida",
            ruta_bd_temporal=fuente,
        )
